from datetime import date
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .models import ReportExportConfig
from .reports import build_balance, build_compte_resultat, build_grand_livre

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
AMOUNT_FORMAT = '#,##0.00;[Red](#,##0.00);-'
CHARCOAL = "404040"
TEXT = "333333"
MUTED = "666666"
ALT_FILL = PatternFill("solid", fgColor="F7F7F7")
TITLE_FILL = PatternFill("solid", fgColor=CHARCOAL)
HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
SUBTITLE_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
SECTION_FILL = PatternFill("solid", fgColor="E7E6E6")
GRID_BORDER = Border(bottom=Side(style="hair", color="D9D9D9"))
HEADER_BORDER = Border(bottom=Side(style="thin", color="A6A6A6"))
TOTAL_BORDER = Border(top=Side(style="thin", color="808080"))


def _export_config():
    return ReportExportConfig.objects.get(key="default").config


def _write_string(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _new_sheet(workbook, name, title, subtitle, headers, details=()):
    sheet = workbook.create_sheet(name)
    sheet.sheet_properties.tabColor = "A6A6A6"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(1, column)
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(vertical="center")
    _write_string(sheet.cell(1, 1), title)
    sheet.cell(1, 1).font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(2, column)
        cell.fill = SUBTITLE_FILL
    _write_string(sheet.cell(2, 1), subtitle)
    sheet.cell(2, 1).font = Font(name="Aptos", size=10, color=MUTED, italic=True)
    sheet.cell(2, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 22
    for index, (label, value) in enumerate(details):
        row_number = 4 + index
        label_cell = sheet.cell(row_number, 1)
        label_cell.value = label
        label_cell.font = Font(name="Aptos", size=10, bold=True, color=CHARCOAL)
        label_cell.fill = SUBTITLE_FILL
        label_cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.merge_cells(
            start_row=row_number,
            start_column=2,
            end_row=row_number,
            end_column=len(headers),
        )
        value_cell = sheet.cell(row_number, 2)
        if isinstance(value, (int, float)):
            value_cell.value = value
            if label == "Solde final":
                value_cell.number_format = AMOUNT_FORMAT
            value_cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            _write_string(value_cell, value or "—")
        value_cell.font = Font(name="Aptos", size=10, color=TEXT)
        value_cell.alignment = Alignment(vertical="center", wrap_text=True)
        value_length = len(str(value or "—"))
        sheet.row_dimensions[row_number].height = max(
            28, 18 * ((value_length + 59) // 60)
        )

    spacer_row = 4 + len(details)
    sheet.row_dimensions[spacer_row].height = 8
    header_row = spacer_row + 1
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column)
        _write_string(cell, value)
        cell.font = Font(name="Aptos", size=10, bold=True, color=CHARCOAL)
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(
            vertical="center",
            horizontal="left" if column == 1 else "center",
            wrap_text=True,
        )
    sheet.row_dimensions[header_row].height = 34
    sheet.freeze_panes = f"A{header_row + 1}"
    return sheet, header_row


def _add_row(sheet, row_number, values, amount_columns=(), kind=None):
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column)
        cell.font = Font(name="Aptos", size=10, color=TEXT)
        cell.border = GRID_BORDER
        cell.alignment = Alignment(vertical="center")
        if row_number % 2 == 0:
            cell.fill = ALT_FILL
        if isinstance(value, date):
            cell.value = value
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif column in amount_columns and value is not None:
            cell.value = value
            cell.number_format = AMOUNT_FORMAT
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            _write_string(cell, value)
        if kind == "section":
            cell.font = Font(name="Aptos", size=10, bold=True, color=CHARCOAL)
            cell.fill = SECTION_FILL
        elif kind == "total":
            cell.font = Font(name="Aptos", size=10, bold=True, color=CHARCOAL)
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_BORDER
    sheet.row_dimensions[row_number].height = 21


def _finish_sheet(sheet, header_row, widths, last_row=None):
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(widths))}{last_row or header_row}"
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.55, bottom=0.55, header=0.2, footer=0.25
    )
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_options.horizontalCentered = True
    sheet.oddHeader.center.text = "&BSmart Compta"
    sheet.oddHeader.center.size = 8
    sheet.oddHeader.center.font = "Arial,Bold"
    sheet.oddFooter.right.text = "Page &P / &N"
    sheet.oddFooter.right.size = 8
    sheet.oddFooter.right.font = "Arial"


def _workbook_bytes():
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.views[0].showHorizontalScroll = True
    workbook.views[0].showVerticalScroll = True
    workbook.properties.creator = "Smart Compta"
    workbook.properties.lastModifiedBy = "Smart Compta"
    workbook.properties.subject = "Rapport financier"
    workbook.properties.keywords = "Comptabilité, DZD, Excel"
    return workbook


def _response(workbook, filename):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), filename


def _subtitle(entreprise, annee):
    return f"{entreprise.nom} | Exercice {annee if annee is not None else 'Tous'} | DZD"


def _company_details(entreprise):
    return [
        ("Désignation de l'entreprise", entreprise.nom),
        ("Activité", entreprise.activite),
        ("Adresse", entreprise.adresse),
        ("N° Article", entreprise.nin),
        ("N° NIF", entreprise.nif),
        ("N° RC", entreprise.nis),
    ]


def _statut_details(annee):
    return [
        ("La période", f"Du : 01/01/{annee} au : 31/12/{annee}"),
        ("Édition", "ÉDITION PROVISOIRE"),
        ("Édition du", timezone.localdate().strftime("%d/%m/%Y")),
    ]


def _format_report_date(value):
    if not value:
        return "—"
    try:
        return date.fromisoformat(value).strftime("%d/%m/%Y")
    except ValueError:
        return value


def _amount_for_prefixes(report, kind, prefixes, configured_prefixes=None):
    key = ",".join(prefixes)
    configured_prefixes = configured_prefixes or {}
    if key in configured_prefixes:
        prefixes = configured_prefixes[key]
    detail = report["detail_charges" if kind == "charges" else "detail_produits"]
    return sum(
        amount for account, amount in detail.items()
        if any(account.startswith(prefix) for prefix in prefixes)
    )


def _sort_income_rows(rows, presentation, config):
    order = config.get("income_order", {}).get(presentation)
    if order:
        positions = {key: index for index, key in enumerate(order)}
        rows.sort(key=lambda row: positions.get(row["key"], len(positions)))
    for item in rows:
        item.pop("key")
    return rows


def _income_statement_rows(report, previous, presentation, config):
    charges, produits = report["charges"], report["produits"]
    previous_charges, previous_produits = previous["charges"], previous["produits"]
    row = lambda label, current, prior, kind="line": {  # noqa: E731
        "key": label,
        "label": config.get("income_labels", {}).get(label, label),
        "current": current, "previous": prior, "kind": kind,
    }
    amount = lambda data, kind, prefixes: _amount_for_prefixes(  # noqa: E731
        data, kind, prefixes, config.get("income_prefixes", {})
    )

    if presentation == "fonction":
        sales = amount(report, "produits", ["70"])
        previous_sales = amount(previous, "produits", ["70"])
        cost = amount(report, "charges", ["60"])
        previous_cost = amount(previous, "charges", ["60"])
        operating = produits["exploitation"] - charges["exploitation"]
        previous_operating = previous_produits["exploitation"] - previous_charges["exploitation"]
        before_tax = operating + produits["financiers"] - charges["financieres"]
        previous_before_tax = (
            previous_operating + previous_produits["financiers"] - previous_charges["financieres"]
        )
        before_exceptional = before_tax - charges["impots"]
        previous_before_exceptional = previous_before_tax - previous_charges["impots"]
        rows = [
            row("Chiffre d'affaires", sales, previous_sales),
            row("Coût des ventes", cost, previous_cost),
            row("MARGE BRUTE", sales - cost, previous_sales - previous_cost, "total"),
            row("Autres produits opérationnels", produits["exploitation"] - sales, previous_produits["exploitation"] - previous_sales),
            row("Coûts commerciaux", amount(report, "charges", ["62"]), amount(previous, "charges", ["62"])),
            row("Charges administratives", amount(report, "charges", ["63", "64"]), amount(previous, "charges", ["63", "64"])),
            row("Autres charges opérationnelles", charges["exploitation"] - cost - amount(report, "charges", ["62", "63", "64"]), previous_charges["exploitation"] - previous_cost - amount(previous, "charges", ["62", "63", "64"])),
            row("RÉSULTAT OPÉRATIONNEL", operating, previous_operating, "total"),
            row("Produits financiers", produits["financiers"], previous_produits["financiers"]),
            row("Charges financières", charges["financieres"], previous_charges["financieres"]),
            row("RÉSULTAT ORDINAIRE AVANT IMPÔT", before_tax, previous_before_tax, "total"),
            row("Impôts exigibles sur les résultats ordinaires", charges["impots"], previous_charges["impots"]),
            row("Impôts différés sur résultats ordinaires", 0, 0),
            row("RÉSULTAT NET DES ACTIVITÉS ORDINAIRES", before_exceptional, previous_before_exceptional, "total"),
            row("Charges extraordinaires", charges["exceptionnelles"], previous_charges["exceptionnelles"]),
            row("Produits extraordinaires", produits["exceptionnels"], previous_produits["exceptionnels"]),
            row("RÉSULTAT NET DE L'EXERCICE", report["resultat"], previous["resultat"], "total"),
            row("Part dans les résultats nets des sociétés mises en équivalence", 0, 0),
            row("RÉSULTAT NET DE L'ENSEMBLE CONSOLIDÉ", report["resultat"], previous["resultat"], "total"),
            row("Dont part des minoritaires", 0, 0),
            row("Part du groupe", report["resultat"], previous["resultat"]),
        ]
        return _sort_income_rows(rows, presentation, config)

    production_prefixes = ["70", "71", "72", "74"]
    sales_cost = amount(report, "charges", ["60"])
    previous_sales_cost = amount(previous, "charges", ["60"])
    services = amount(report, "charges", ["61", "62"])
    previous_services = amount(previous, "charges", ["61", "62"])
    personnel = amount(report, "charges", ["63"])
    previous_personnel = amount(previous, "charges", ["63"])
    taxes = amount(report, "charges", ["64"])
    previous_taxes = amount(previous, "charges", ["64"])
    production = amount(report, "produits", production_prefixes)
    previous_production = amount(previous, "produits", production_prefixes)
    consumption = sales_cost + services
    previous_consumption = previous_sales_cost + previous_services
    value_added = production - consumption
    previous_value_added = previous_production - previous_consumption
    gross_surplus = value_added - personnel - taxes
    previous_gross_surplus = previous_value_added - previous_personnel - previous_taxes
    operating = produits["exploitation"] - charges["exploitation"]
    previous_operating = previous_produits["exploitation"] - previous_charges["exploitation"]
    financial = produits["financiers"] - charges["financieres"]
    previous_financial = previous_produits["financiers"] - previous_charges["financieres"]
    before_tax = operating + financial
    previous_before_tax = previous_operating + previous_financial
    ordinary_result = before_tax - charges["impots"]
    previous_ordinary_result = previous_before_tax - previous_charges["impots"]
    rows = [
        row("Chiffre d'affaires", amount(report, "produits", ["70"]), amount(previous, "produits", ["70"])),
        row("Variation stocks produits finis et en cours", amount(report, "produits", ["71"]), amount(previous, "produits", ["71"])),
        row("Production immobilisée", amount(report, "produits", ["72"]), amount(previous, "produits", ["72"])),
        row("Subventions d'exploitation", amount(report, "produits", ["74"]), amount(previous, "produits", ["74"])),
        row("I - PRODUCTION DE L'EXERCICE", production, previous_production, "total"),
        row("Achats consommés", sales_cost, previous_sales_cost),
        row("Services extérieurs et autres services", services, previous_services),
        row("II - CONSOMMATION DE L'EXERCICE", consumption, previous_consumption, "total"),
        row("III - VALEUR AJOUTÉE D'EXPLOITATION (I - II)", value_added, previous_value_added, "total"),
        row("Charges de personnel", personnel, previous_personnel),
        row("Impôts, taxes et versements assimilés", taxes, previous_taxes),
        row("IV - EXCÉDENT BRUT D'EXPLOITATION", gross_surplus, previous_gross_surplus, "total"),
        row("Autres produits opérationnels", amount(report, "produits", ["75"]), amount(previous, "produits", ["75"])),
        row("Autres charges opérationnelles", amount(report, "charges", ["65"]), amount(previous, "charges", ["65"])),
        row("Dotations aux amortissements et aux provisions", amount(report, "charges", ["68"]), amount(previous, "charges", ["68"])),
        row("Reprise sur pertes de valeur et provisions", amount(report, "produits", ["78"]), amount(previous, "produits", ["78"])),
        row("V - RÉSULTAT OPÉRATIONNEL", operating, previous_operating, "total"),
        row("Produits financiers", produits["financiers"], previous_produits["financiers"]),
        row("Charges financières", charges["financieres"], previous_charges["financieres"]),
        row("VI - RÉSULTAT FINANCIER", financial, previous_financial, "total"),
        row("VII - RÉSULTAT ORDINAIRE AVANT IMPÔTS (V + VI)", before_tax, previous_before_tax, "total"),
        row("Impôts exigibles sur les résultats ordinaires", charges["impots"], previous_charges["impots"]),
        row("Impôts différés", 0, 0),
        row("TOTAL DES PRODUITS DES ACTIVITÉS ORDINAIRES", produits["exploitation"] + produits["financiers"], previous_produits["exploitation"] + previous_produits["financiers"], "total"),
        row("TOTAL DES CHARGES DES ACTIVITÉS ORDINAIRES", charges["exploitation"] + charges["financieres"] + charges["impots"], previous_charges["exploitation"] + previous_charges["financieres"] + previous_charges["impots"], "total"),
        row("VIII - RÉSULTAT NET DES ACTIVITÉS ORDINAIRES", ordinary_result, previous_ordinary_result, "total"),
        row("Éléments extraordinaires (produits)", produits["exceptionnels"], previous_produits["exceptionnels"]),
        row("Éléments extraordinaires (charges)", charges["exceptionnelles"], previous_charges["exceptionnelles"]),
        row("IX - RÉSULTAT EXTRAORDINAIRE", produits["exceptionnels"] - charges["exceptionnelles"], previous_produits["exceptionnels"] - previous_charges["exceptionnelles"], "total"),
        row("X - RÉSULTAT NET DE L'EXERCICE", report["resultat"], previous["resultat"], "total"),
        row("Part dans les résultats nets des sociétés mises en équivalence", 0, 0),
        row("XI - RÉSULTAT NET DE L'ENSEMBLE CONSOLIDÉ", report["resultat"], previous["resultat"], "total"),
        row("Dont part des minoritaires", 0, 0),
        row("Part du groupe", report["resultat"], previous["resultat"]),
    ]
    return _sort_income_rows(rows, presentation, config)


def _balance_accounts(data):
    return {
        account["compte"]: account
        for group in data["classes"]
        for account in group["comptes"]
    }


def _matches(account, prefixes, exclusions=()):
    return any(account.startswith(prefix) for prefix in prefixes) and not any(
        account.startswith(prefix) for prefix in exclusions
    )


def _asset_rows(current, previous, specs):
    rows = []
    for spec in specs:
        values = {"brut": 0, "amortissement": 0, "net": 0, "previous": 0}
        for code in set(current) | set(previous):
            account = current.get(code) or previous.get(code)
            if _matches(account["compte"], spec["prefixes"], spec.get("exclusions", ())):
                now, prior = current.get(code, {}), previous.get(code, {})
                values["brut"] += now.get("solde_debiteur", 0)
                values["net"] += now.get("solde_debiteur", 0) - now.get("solde_crediteur", 0)
                values["previous"] += prior.get("solde_debiteur", 0) - prior.get("solde_crediteur", 0)
            if _matches(account["compte"], spec.get("provisionPrefixes", ()), spec.get("provisionExclusions", ())):
                now, prior = current.get(code, {}), previous.get(code, {})
                values["amortissement"] += now.get("solde_crediteur", 0)
                values["net"] -= now.get("solde_crediteur", 0) - now.get("solde_debiteur", 0)
                values["previous"] -= prior.get("solde_crediteur", 0) - prior.get("solde_debiteur", 0)
        rows.append({**spec, **values, "kind": "group" if not spec["prefixes"] else "line"})
    return rows


def _liability_rows(current, previous, specs):
    rows = []
    for spec in specs:
        amounts = []
        for account_map in (current, previous):
            amount = 0
            for code, account in account_map.items():
                if _matches(code, spec["prefixes"], spec.get("exclusions", ())):
                    credit, debit = account["solde_crediteur"], account["solde_debiteur"]
                    amount += credit - debit if spec.get("signed") else credit
            amounts.append(amount)
        rows.append({**spec, "kind": "line", "net": amounts[0], "previous": amounts[1]})
    return rows


def _sum_asset_rows(rows):
    return {
        key: sum(row.get(key, 0) for row in rows)
        for key in ("brut", "amortissement", "net", "previous")
    }


def _sum_liability_rows(rows):
    return {
        "current": sum(row.get("net", 0) for row in rows),
        "previous": sum(row.get("previous", 0) for row in rows),
    }


def _bilan_data(entreprise, year, config):
    current_data = build_balance(entreprise, year)
    previous_data = build_balance(entreprise, year - 1)
    current, previous = _balance_accounts(current_data), _balance_accounts(previous_data)
    layout = config["bilan"]

    non_current = _asset_rows(current, previous, layout["non_current_assets"])
    current_assets = _asset_rows(current, previous, layout["current_assets"])
    non_current_totals = _sum_asset_rows(non_current)
    current_totals = _sum_asset_rows(current_assets)
    actif_rows = (
        [{"label": layout["labels"]["ACTIF NON COURANTS"], "kind": "section"}]
        + non_current
        + [{"label": layout["labels"]["TOTAL ACTIF NON COURANT"], "kind": "total", **non_current_totals}]
        + [{"label": layout["labels"]["ACTIF COURANT"], "kind": "section"}]
        + current_assets
        + [{"label": layout["labels"]["TOTAL ACTIF COURANT"], "kind": "total", **current_totals}]
    )
    actif_totals = {
        key: non_current_totals[key] + current_totals[key]
        for key in ("brut", "amortissement", "net", "previous")
    }

    equity = _liability_rows(current, previous, layout["equity"])
    non_current_liabilities = _liability_rows(current, previous, layout["non_current_liabilities"])
    current_liabilities = _liability_rows(current, previous, layout["current_liabilities"])
    equity_totals = _sum_liability_rows(equity)
    non_current_liability_totals = _sum_liability_rows(non_current_liabilities)
    current_liability_totals = _sum_liability_rows(current_liabilities)
    passif_rows = (
        [{"label": layout["labels"]["CAPITAUX PROPRES"], "kind": "section"}] + equity
        + [{"label": layout["labels"]["TOTAL I"], "kind": "total", "net": equity_totals["current"], "previous": equity_totals["previous"]}]
        + [{"label": layout["labels"]["PASSIFS NON COURANTS"], "kind": "section"}] + non_current_liabilities
        + [{"label": layout["labels"]["TOTAL PASSIFS NON COURANTS II"], "kind": "total", "net": non_current_liability_totals["current"], "previous": non_current_liability_totals["previous"]}]
        + [{"label": layout["labels"]["PASSIFS COURANTS"], "kind": "section"}] + current_liabilities
        + [{"label": layout["labels"]["TOTAL PASSIFS COURANTS III"], "kind": "total", "net": current_liability_totals["current"], "previous": current_liability_totals["previous"]}]
    )
    passif_totals = {
        "current": equity_totals["current"] + non_current_liability_totals["current"] + current_liability_totals["current"],
        "previous": equity_totals["previous"] + non_current_liability_totals["previous"] + current_liability_totals["previous"],
    }
    return actif_rows, actif_totals, passif_rows, passif_totals


def export_balance(entreprise, annee):
    data = build_balance(entreprise, annee)
    workbook = _workbook_bytes()
    sheet, header = _new_sheet(
        workbook, "Balance", "Balance", _subtitle(entreprise, annee),
        ["Compte", "Libellé", "Débit (DZD)", "Crédit (DZD)", "Solde débiteur (DZD)", "Solde créditeur (DZD)"],
        [("Entreprise", entreprise.nom), ("Exercice", annee or "Tous")],
    )
    row = header + 1
    for group in data["classes"]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        _add_row(sheet, row, [f"Classe {group['classe']} — {group['label']}"], kind="section")
        row += 1
        for account in group["comptes"]:
            _add_row(sheet, row, [
                account["compte"], account["libelle"], account["debit"], account["credit"],
                account["solde_debiteur"], account["solde_crediteur"],
            ], (3, 4, 5, 6))
            row += 1
    totals = data["totals"]
    _add_row(sheet, row, [
        "TOTAL", "", totals["debit"], totals["credit"],
        totals["solde_debiteur"], totals["solde_crediteur"],
    ], (3, 4, 5, 6), "total")
    _finish_sheet(sheet, header, [20, 48, 22, 22, 26, 26], row)
    return _response(workbook, f"balance-{annee or 'tous'}.xlsx")


def export_compte_resultat(entreprise, annee, presentation):
    if presentation not in ("fonction", "nature"):
        presentation = "fonction"
    current = build_compte_resultat(entreprise, annee)
    previous = build_compte_resultat(entreprise, annee - 1) if annee is not None else current
    config = _export_config()
    rows = _income_statement_rows(current, previous, presentation, config)
    workbook = _workbook_bytes()
    sheet, header = _new_sheet(
        workbook, "Compte de résultat", f"Compte de résultat par {presentation}",
        _subtitle(entreprise, annee),
        ["Rubriques", "Note", f"Exercice {annee or 'courant'} (DZD)", f"Exercice {(annee - 1) if annee else 'précédent'} (DZD)"],
        _company_details(entreprise) + _statut_details(annee),
    )
    for index, item in enumerate(rows, start=header + 1):
        _add_row(
            sheet, index, [item["label"], "", item["current"], item["previous"]],
            (3, 4), item["kind"],
        )
    _finish_sheet(sheet, header, [62, 12, 24, 24], header + len(rows))
    return _response(workbook, f"compte-de-resultat-{annee or 'tous'}-{presentation}.xlsx")


def export_grand_livre(entreprise, annee, start, end, account_number):
    if not account_number:
        raise ValueError("Sélectionnez un compte avant d’exporter le grand livre.")
    data = build_grand_livre(entreprise, annee, start, end)
    account = next((item for item in data["comptes"] if item["compte"] == account_number), None)
    if account is None:
        raise LookupError("Le compte sélectionné ne contient aucun mouvement pour cette période.")
    workbook = _workbook_bytes()
    sheet, header = _new_sheet(
        workbook, "Grand livre", f"Grand livre — {account['compte']} {account['libelle']}",
        f"{_subtitle(entreprise, annee)} | Du {_format_report_date(start)} au {_format_report_date(end)}",
        ["Date", "N° pièce", "Libellé", "Débit (DZD)", "Crédit (DZD)", "Solde (DZD)"],
        [
            ("Entreprise", entreprise.nom),
            ("Exercice", annee or "Tous"),
            (
                "Période",
                f"Du {start and _format_report_date(start) or 'début'} "
                f"au {end and _format_report_date(end) or 'fin'}",
            ),
            ("Compte", f"{account['compte']} — {account['libelle']}"),
            ("Solde final", account["solde_final"]),
        ],
    )
    for index, movement in enumerate(account["mouvements"], start=header + 1):
        _add_row(sheet, index, [
            date.fromisoformat(movement["date"]), movement["numero_piece"], movement["libelle"],
            movement["debit"], movement["credit"], movement["solde"],
        ], (4, 5, 6))
    _add_row(
        sheet, header + len(account["mouvements"]) + 1,
        ["", "", "SOLDE FINAL", "", "", account["solde_final"]],
        (6,), "total",
    )
    _finish_sheet(sheet, header, [16, 20, 48, 20, 20, 22], header + len(account["mouvements"]) + 1)
    return _response(workbook, f"grand-livre-{account['compte']}-{annee or 'tous'}.xlsx")


def export_bilan(entreprise, annee):
    current_year = annee
    config = _export_config()["bilan"]
    labels = config["labels"]
    actif_rows, actif_totals, passif_rows, passif_totals = _bilan_data(
        entreprise, current_year, {"bilan": config}
    )
    workbook = _workbook_bytes()
    subtitle = _subtitle(entreprise, current_year)
    sheet, header = _new_sheet(
        workbook, "Actif", "Bilan — Actif", subtitle,
        ["Rubrique", "Note", f"Exercice {current_year} — Brut (DZD)", f"Exercice {current_year} — Amort.-Prov. (DZD)", f"Exercice {current_year} — Net (DZD)", f"Exercice {current_year - 1} — Net (DZD)"],
        _company_details(entreprise) + _statut_details(current_year),
    )
    for index, item in enumerate(actif_rows, start=header + 1):
        _add_row(sheet, index, [
            f"  {item['label']}" if item.get("indent") else item["label"], "",
            item.get("brut"), item.get("amortissement"), item.get("net"), item.get("previous"),
        ], (3, 4, 5, 6), item["kind"])
    row = header + len(actif_rows) + 1
    _add_row(sheet, row, [
        labels["TOTAL GÉNÉRAL ACTIF"], "", actif_totals["brut"], actif_totals["amortissement"],
        actif_totals["net"], actif_totals["previous"],
    ], (3, 4, 5, 6), "total")
    _finish_sheet(sheet, header, [55, 12, 25, 28, 24, 24], row)

    sheet, header = _new_sheet(
        workbook, "Passif", "Bilan — Passif", subtitle,
        ["Rubrique", "Note", f"Exercice {current_year} (DZD)", f"Exercice {current_year - 1} (DZD)"],
        _company_details(entreprise) + _statut_details(current_year),
    )
    for index, item in enumerate(passif_rows, start=header + 1):
        _add_row(sheet, index, [
            f"  {item['label']}" if item.get("indent") else item["label"], "",
            item.get("net"), item.get("previous"),
        ], (3, 4), item["kind"])
    row = header + len(passif_rows) + 1
    _add_row(sheet, row, [
        labels["TOTAL GÉNÉRAL PASSIF"], "", passif_totals["current"], passif_totals["previous"],
    ], (3, 4), "total")
    _finish_sheet(sheet, header, [58, 12, 25, 25], row)
    return _response(workbook, f"bilan-{current_year}.xlsx")
